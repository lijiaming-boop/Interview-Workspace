from copy import deepcopy
from pathlib import Path

from docx import Document


SOURCE = Path(r"E:\contion simulation\李佳明_简历_Agent开发岗位_实习个人项目版.docx")
OUTPUT = Path(r"E:\contion simulation\李佳明_简历_Agent开发岗位_暖通工具测算版.docx")


def replace_paragraph(paragraph, text: str) -> None:
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


doc = Document(str(SOURCE))
replace_paragraph(
    doc.paragraphs[16],
    "暖通工具测算项目：参与制冷站历史工况数据清洗、缓存化和字段统一，建立冷却塔散热、冷却泵功率、主机功率与逼近温差子模型；采用 70%/30% 划分和时间分块交叉验证评估泛化。",
)
replace_paragraph(
    doc.paragraphs[17],
    "工具测算与闭环联合寻优：将实测主机功率 ChPower、冷却泵功率 CWPPower 纳入目标函数，完成塔组合枚举 + SLSQP 连续调频；加入流量/风量包络、温差、出水温度与单塔容量约束，输出泵/塔频率和净功率建议。",
)
doc.save(str(OUTPUT))
print(OUTPUT)
