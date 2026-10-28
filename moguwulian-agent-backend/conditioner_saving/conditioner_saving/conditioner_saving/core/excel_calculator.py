"""Excel 计算引擎

使用 formulas 库执行 Excel 公式计算（纯 Python 实现，无需外部依赖）
"""
import json
import shutil
from pathlib import Path
from datetime import datetime
from typing import Dict, Any, Optional
from config.cell_mapping import CellMapping
from config.template_config import TemplateConfig
from config import WeatherService

# 计算引擎版本（P1 改造：结果与公式模板/单元格映射版本绑定，便于追溯口径变化）
CALCULATION_ENGINE_VERSION = "2.0.0"


class ExcelCalculator:
    """Excel 计算引擎 - 数据回填和结果提取"""

    def __init__(self, template_path: Optional[str] = None, weather_service: Optional[WeatherService] = None):
        """初始化计算器

        Args:
            template_path: Excel 模板文件路径（可选，如果不提供则根据 station_type 动态选择）
            weather_service: 气象服务（可选，用于冷却塔计算）
        """
        self.template_path = template_path
        self.mapper = CellMapping()
        self.template_config = TemplateConfig()
        self.weather_service = weather_service or WeatherService()

        project_root = Path(__file__).resolve().parent.parent
        self.template_dir = project_root / "00_sample" / "模板测算表"
        self.quotation_template_path = self.template_dir / "2026年空调报价模板.xlsx"
        self.scheme_template_path = self.template_dir / "空调系统数智化升级方案模板.docx"
        self.quotation_filename = "ac_quotation.xlsx"
        self.scheme_filename = "ac_scheme.docx"

    def calculate(
        self,
        filled_data_path: str,
        output_path: str
    ) -> str:
        """执行节能率计算

        Args:
            filled_data_path: 填充后的数据 JSON 路径
            output_path: 结果输出 JSON 路径

        Returns:
            结果文件路径

        Raises:
            ValueError: 站房类型为 unsupported
        """
        # 读取填充后的数据
        with open(filled_data_path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        # 如果没有指定模板路径，根据 station_type 自动选择
        if self.template_path is None:
            station_type = data.get("station_type", "standard")

            # 检查是否为不支持的类型
            if station_type == "unsupported":
                raise ValueError(
                    "不支持的站房类型。该类型站房（如包含热水泵、水蓄冷、中低温机组）"
                    "不在自动节能评估范围内。"
                )

            # 动态获取模板路径
            template_path = self.template_config.get_template_path(station_type)
        else:
            template_path = self.template_path

        # 创建临时 Excel 文件
        temp_excel = Path(output_path).parent / "calculation_result.xlsx"
        shutil.copy(template_path, temp_excel)

        # 公式计算失败必须显式报错，严禁静默降级到 mock/默认值——
        # 前端展示的是给销售报价用的节能率，假数据比报错危害更大
        result = self._calculate_with_formulas(data, temp_excel)

        # 保存结果
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(result, f, ensure_ascii=False, indent=2)

        # 生成简化版Excel文件
        self._create_simplified_excel(temp_excel, Path(output_path).parent / "simplified_result.xlsx")
        
        return str(output_path)
    
    def _create_simplified_excel(self, original_excel_path: Path, simplified_excel_path: Path):
        """创建简化版Excel文件，创建新工作表并只复制需要保留的行
        
        Args:
            original_excel_path: 原始Excel文件路径
            simplified_excel_path: 简化版Excel文件路径
        """
        import openpyxl
        
        # 加载原始工作簿
        wb = openpyxl.load_workbook(original_excel_path)
        
        # 定义要保留的行范围（不删除的行）
        # 保留第1-2行，7-8行，11-12行，16-17行，21-22行，25行及以后
        rows_to_keep = set(list(range(1, 3)) + list(range(7, 9)) + list(range(11, 13)) + list(range(16, 18)) + list(range(21, 23)) + list(range(25, 1000)))
        
        # 处理上限sheet
        if "【1】按冷站电量数据-上限" in wb.sheetnames:
            self._copy_filtered_sheet(wb, "【1】按冷站电量数据-上限", rows_to_keep)
        
        # 处理下限sheet
        if "【2】按冷站电量数据-下限" in wb.sheetnames:
            self._copy_filtered_sheet(wb, "【2】按冷站电量数据-下限", rows_to_keep)
        
        # 删除气象数据sheet
        if "全国主要城市气象情况" in wb.sheetnames:
            del wb["全国主要城市气象情况"]
        
        # 保存简化版Excel
        wb.save(simplified_excel_path)
        wb.close()
    
    def _copy_filtered_sheet(self, wb, sheet_name: str, rows_to_keep: set):
        """复制工作表，只保留指定的行
        
        Args:
            wb: 工作簿对象
            sheet_name: 工作表名称
            rows_to_keep: 要保留的行号集合
        """
        original_sheet = wb[sheet_name]
        
        # 创建新工作表
        new_sheet_name = sheet_name
        new_sheet = wb.create_sheet(new_sheet_name)
        
        # 记录原始合并单元格
        merged_cells_info = []
        for merged_range in original_sheet.merged_cells.ranges:
            start_row, end_row = merged_range.min_row, merged_range.max_row
            start_col, end_col = merged_range.min_col, merged_range.max_col
            merged_cells_info.append({
                'start_row': start_row,
                'end_row': end_row,
                'start_col': start_col,
                'end_col': end_col
            })
        
        # 创建行号映射（原始行号 -> 新行号）
        row_mapping = {}
        new_row = 1
        for original_row in range(1, original_sheet.max_row + 1):
            if original_row in rows_to_keep:
                row_mapping[original_row] = new_row
                new_row += 1
        
        # 复制单元格数据
        for original_row in range(1, original_sheet.max_row + 1):
            if original_row in rows_to_keep:
                new_row = row_mapping[original_row]
                for original_col in range(1, original_sheet.max_column + 1):
                    original_cell = original_sheet.cell(row=original_row, column=original_col)
                    new_cell = new_sheet.cell(row=new_row, column=original_col)
                    
                    # 复制值（此时公式已经被替换为计算结果）
                    new_cell.value = original_cell.value
                    
                    # 复制字体样式
                    if original_cell.has_style:
                        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Protection
                        
                        # 复制字体
                        if original_cell.font:
                            from openpyxl.styles.colors import Color
                            
                            # 处理颜色对象
                            font_color = None
                            if original_cell.font.color:
                                if hasattr(original_cell.font.color, 'rgb'):
                                    # 如果是RGB对象，直接使用，不需要转换
                                    font_color = original_cell.font.color
                                elif hasattr(original_cell.font.color, 'theme'):
                                    # 如果是主题颜色
                                    font_color = Color(theme=original_cell.font.color.theme)
                                elif hasattr(original_cell.font.color, 'indexed'):
                                    # 如果是索引颜色
                                    font_color = Color(indexed=original_cell.font.color.indexed)
                            
                            new_cell.font = Font(
                                name=original_cell.font.name,
                                size=original_cell.font.size,
                                bold=original_cell.font.bold,
                                italic=original_cell.font.italic,
                                underline=original_cell.font.underline,
                                strike=original_cell.font.strike,
                                color=font_color
                            )
                        
                        # 复制填充（背景色）
                        if original_cell.fill and hasattr(original_cell.fill, 'fgColor'):
                            from openpyxl.styles.colors import Color
                            
                            # 处理前景色
                            fg_color = None
                            if original_cell.fill.fgColor:
                                if hasattr(original_cell.fill.fgColor, 'rgb'):
                                    # 直接使用RGB对象，不需要转换
                                    fg_color = original_cell.fill.fgColor
                                elif hasattr(original_cell.fill.fgColor, 'theme'):
                                    fg_color = Color(theme=original_cell.fill.fgColor.theme)
                                elif hasattr(original_cell.fill.fgColor, 'indexed'):
                                    fg_color = Color(indexed=original_cell.fill.fgColor.indexed)
                            
                            new_cell.fill = PatternFill(
                                fill_type=original_cell.fill.fill_type,
                                start_color=fg_color,
                                end_color=fg_color
                            )
                        
                        # 复制边框
                        if original_cell.border:
                            from openpyxl.styles.colors import Color
                            
                            def copy_side(side):
                                if side:
                                    # 处理边框颜色
                                    side_color = None
                                    if side.color:
                                        if hasattr(side.color, 'rgb'):
                                            # 直接使用RGB对象，不需要转换
                                            side_color = side.color
                                        elif hasattr(side.color, 'theme'):
                                            side_color = Color(theme=side.color.theme)
                                        elif hasattr(side.color, 'indexed'):
                                            side_color = Color(indexed=side.color.indexed)
                                    
                                    return Side(
                                        style=side.style,
                                        color=side_color
                                    )
                                return None
                            
                            new_cell.border = Border(
                                left=copy_side(original_cell.border.left),
                                right=copy_side(original_cell.border.right),
                                top=copy_side(original_cell.border.top),
                                bottom=copy_side(original_cell.border.bottom),
                                diagonal=copy_side(original_cell.border.diagonal),
                                diagonal_direction=original_cell.border.diagonal_direction,
                                outline=original_cell.border.outline,
                                start=copy_side(original_cell.border.start),
                                end=copy_side(original_cell.border.end)
                            )
                        
                        # 复制对齐方式
                        if original_cell.alignment:
                            new_cell.alignment = Alignment(
                                horizontal=original_cell.alignment.horizontal,
                                vertical=original_cell.alignment.vertical,
                                text_rotation=original_cell.alignment.text_rotation,
                                wrap_text=original_cell.alignment.wrap_text,
                                shrink_to_fit=original_cell.alignment.shrink_to_fit,
                                indent=original_cell.alignment.indent
                            )
                        
                        # 复制保护
                        if original_cell.protection:
                            new_cell.protection = Protection(
                                locked=original_cell.protection.locked,
                                hidden=original_cell.protection.hidden
                            )
                    
                    # 复制数字格式
                    new_cell.number_format = original_cell.number_format
        
        # 复制列宽
        for col_idx, column_dim in original_sheet.column_dimensions.items():
            new_sheet.column_dimensions[col_idx].width = column_dim.width
        
        # 复制行高
        for row_idx, row_dim in original_sheet.row_dimensions.items():
            if row_idx in row_mapping:
                new_sheet.row_dimensions[row_mapping[row_idx]].height = row_dim.height
        
        # 应用合并单元格
        for merged_info in merged_cells_info:
            original_start_row = merged_info['start_row']
            original_end_row = merged_info['end_row']
            
            # 检查合并单元格是否完全在保留的行范围内
            if original_start_row in row_mapping and original_end_row in row_mapping:
                new_start_row = row_mapping[original_start_row]
                new_end_row = row_mapping[original_end_row]
                new_start_col = merged_info['start_col']
                new_end_col = merged_info['end_col']
                
                try:
                    new_sheet.merge_cells(
                        start_row=new_start_row,
                        start_column=new_start_col,
                        end_row=new_end_row,
                        end_column=new_end_col
                    )
                except Exception:
                    pass
        
        # 删除原始工作表
        del wb[sheet_name]

    def _calculate_with_formulas(
        self,
        data: Dict[str, Any],
        excel_path: Path
    ) -> Dict[str, Any]:
        """使用 formulas 库执行计算

        计算上限和下限两个sheet，返回节能率范围

        Args:
            data: 填充后的数据字典
            excel_path: Excel 文件路径

        Returns:
            计算结果字典（包含范围）
        """
        try:
            import formulas  # noqa: F401  可用性探测
            import openpyxl  # noqa: F401
        except ImportError as e:
            raise RuntimeError(f"Required library not installed: {e}")

        # 计算下限（保守估计）
        result_lower = self._calculate_single_sheet(data, excel_path, "lower")

        # 计算上限（理想情况）
        result_upper = self._calculate_single_sheet(data, excel_path, "upper")

        # 合并结果，构建范围
        result = {
            # 节能率范围
            "energy_saving_rate_min": result_lower["saving_rate"],
            "energy_saving_rate_max": result_upper["saving_rate"],
            "energy_saving_rate_range": (
                f"{result_lower['saving_rate']*100:.1f}% - "
                f"{result_upper['saving_rate']*100:.1f}%"
            ),

            # 向后兼容：保留原有单一值字段（使用下限）
            "energy_saving_rate": result_lower["saving_rate"],

            # 分项节能量范围
            "chiller_saving_min": result_lower["chiller_saving"],
            "chiller_saving_max": result_upper["chiller_saving"],
            "chiller_saving": result_lower["chiller_saving"],  # 向后兼容

            "pump_saving_min": result_lower["pump_saving"],
            "pump_saving_max": result_upper["pump_saving"],
            "pump_saving": result_lower["pump_saving"],  # 向后兼容

            "tower_saving_min": result_lower["tower_saving"],
            "tower_saving_max": result_upper["tower_saving"],
            "tower_saving": result_lower["tower_saving"],  # 向后兼容

            "total_saving_min": result_lower["total_saving"],
            "total_saving_max": result_upper["total_saving"],
            "total_saving": result_lower["total_saving"],  # 向后兼容

            # 元数据
            "completeness": data["_metadata"]["final_completeness"],
            "filled_count": data["_metadata"]["fill_count"],
            "calculation_timestamp": datetime.now().isoformat(),
            "calculation_method": "dual_sheet",  # 标记使用双sheet计算

            # P1 改造：版本绑定 + 单位口径显式化 + 原始总能耗
            "engine_version": CALCULATION_ENGINE_VERSION,
            "cell_mapping_version": self.mapper.MAPPING_VERSION,
            "energy_unit": "万kWh",
            "total_energy_original": round(sum(
                (m.get("total_energy") or 0) for m in data.get("monthly_data", []) or []
            ), 4),
            "filling_defaults_version": (data.get("_metadata") or {}).get("defaults_version", "unversioned"),
        }

        self._validate_result(result, data)

        return result

    def _validate_result(self, result: Dict[str, Any], data: Dict[str, Any]) -> None:
        """计算结果合理性校验（P1 改造：失败关闭，杜绝把故障伪装成"无节能效果"）

        检查项（改进方案第五节）：
        1. 数值有限性：所有指标必须是有限数（拒绝 NaN/Inf）
        2. 节能率上下限顺序：min <= max
        3. 节能率合理范围：[-50%, 100%] 之外视为公式/映射故障
        4. 分项与总量一致性：冷机+水泵+冷却塔 与 总量 偏差超 2% 报错
        5. 节能量不得超过原始能耗
        """
        import math

        numeric_fields = (
            "energy_saving_rate_min", "energy_saving_rate_max",
            "chiller_saving", "pump_saving", "tower_saving", "total_saving",
        )
        for f in numeric_fields:
            v = result.get(f)
            if v is None or not isinstance(v, (int, float)) or not math.isfinite(v):
                raise RuntimeError(
                    f"计算结果字段 {f} 不是有限数值（{v!r}）——"
                    "公式计算或单元格映射异常，拒绝输出可能误导报价的结果。"
                )

        rate_min = result["energy_saving_rate_min"]
        rate_max = result["energy_saving_rate_max"]
        if rate_min > rate_max:
            raise RuntimeError(
                f"节能率下限({rate_min:.2%})大于上限({rate_max:.2%})——"
                "上限/下限 sheet 结果交叉，请检查模板与单元格映射。"
            )
        if not (-0.5 <= rate_min <= 1.0 and -0.5 <= rate_max <= 1.0):
            raise RuntimeError(
                f"节能率范围 [{rate_min:.2%}, {rate_max:.2%}] 超出合理区间 [-50%, 100%]"
                "——通常是输入数据或公式映射故障，拒绝输出。"
            )

        parts = result["chiller_saving"] + result["pump_saving"] + result["tower_saving"]
        total = result["total_saving"]
        if total != 0 and abs(parts - total) > max(abs(total) * 0.02, 0.01):
            raise RuntimeError(
                f"分项节能量合计({parts:.2f})与总节能量({total:.2f})偏差超过 2%——"
                "公式模板或结果单元格映射可能失配，拒绝输出。"
            )

        original = result.get("total_energy_original") or 0
        if original > 0 and total > original:
            raise RuntimeError(
                f"总节能量({total:.2f} 万kWh)超过原始总能耗({original:.2f} 万kWh)——"
                "物理上不可能，公式或输入数据存在故障，拒绝输出。"
            )
    
    def format_export_results(self, data: Dict[str, Any], calculation_result: Dict[str, Any]) -> Dict[str, Any]:
        """格式化导出结果（单位：万kWh）

        P1 正确性修复（改进方案第五节）：
        1. 原实现 total_energy_original 误用了 total_saving——"原始总能耗"
           展示的是节能量；
        2. 原实现对已是"万kWh"的数值再除以 10000（模板引擎内部单位统一为
           万kWh：月度电量 ~30、年节能量 ~41），导致摘要全部缩小一万倍。

        现统一口径：内部数值即万kWh，展示不再换算；原始总能耗 = 各月电量之和。
        """
        monthly_rows = data.get("monthly_data", []) or []
        total_original = sum(m.get("total_energy") or 0 for m in monthly_rows)
        total_saving = calculation_result.get("total_saving", 0) or 0

        formatted_result = {
            "project_info": {
                "project_region": data.get("project_region"),
                "station_type": data.get("station_type"),
                "completeness": data.get("_metadata", {}).get("final_completeness")
            },
            "unit": "万kWh",
            "monthly_data": [],
            "summary": {
                "total_energy_original": round(total_original, 2),
                "energy_saving_amount": round(total_saving, 2),
                "energy_saving_rate_min": calculation_result.get("energy_saving_rate_min"),
                "energy_saving_rate_max": calculation_result.get("energy_saving_rate_max"),
                "energy_saving_rate_range": calculation_result.get("energy_saving_rate_range")
            }
        }

        # 格式化月度数据，只保留原始电量和节能量（单位：万kWh，不换算）
        saving_rate = calculation_result.get("energy_saving_rate", 0) or 0
        for month_data in monthly_rows:
            original = month_data.get("total_energy") or 0
            formatted_month = {
                "month": month_data.get("month"),
                "original_energy": round(original, 2),
                "energy_saving": round(original * saving_rate, 2)
            }
            formatted_result["monthly_data"].append(formatted_month)

        return formatted_result

    def _generate_quotation(self, session_dir: str, filled_data: dict, result: dict):
        """生成空调报价单

        在产品价目表模板最前面插入「项目节能测算」汇总表，
        写入项目信息与节能计算结果，另存到会话目录。
        """
        if not self.quotation_template_path.exists():
            print(f"报价模板不存在: {self.quotation_template_path}")
            return None

        from openpyxl import load_workbook

        quotation_path = Path(session_dir) / self.quotation_filename

        try:
            wb = load_workbook(self.quotation_template_path)

            ws = wb.create_sheet("项目节能测算", 0)
            ws.column_dimensions["A"].width = 22
            ws.column_dimensions["B"].width = 30

            rows = [
                ("项目区域", filled_data.get("project_region", "")),
                ("站房类型", filled_data.get("station_type", "")),
                ("", ""),
                ("电价 (元/kWh)", filled_data.get("electricity_price", "")),
                ("数据完整率", f"{result.get('completeness', 0) * 100:.1f}%"),
                ("", ""),
                ("总节电量 (万kWh/年)", round(float(result.get("total_saving", 0)), 2)),
                ("冷机节电 (万kWh/年)", round(float(result.get("chiller_saving", 0)), 2)),
                ("水泵节电 (万kWh/年)", round(float(result.get("pump_saving", 0)), 2)),
                ("冷却塔节电 (万kWh/年)", round(float(result.get("tower_saving", 0)), 2)),
                ("", ""),
                ("综合节能率", f"{result.get('energy_saving_rate', 0) * 100:.1f}%"),
                ("节能率范围", result.get("energy_saving_rate_range", "")),
            ]
            ws["A1"] = "空调系统节能测算汇总"
            for idx, (label, value) in enumerate(rows, start=2):
                ws.cell(row=idx, column=1, value=label)
                ws.cell(row=idx, column=2, value=value)

            wb.save(quotation_path)
            print(f"报价单已生成: {quotation_path}")
            return str(quotation_path)

        except Exception as e:
            print(f"生成报价单失败: {str(e)}")
            return None

    def _generate_scheme_document(self, session_dir: str, filled_data: dict, result: dict):
        """生成空调系统数智化升级方案文档

        基于 docx 模板将正文段落中的【】占位符替换为实际项目数据，
        固定输出到会话目录，供前端下载。
        """
        if not self.scheme_template_path.exists():
            print(f"方案模板不存在: {self.scheme_template_path}")
            return None

        scheme_path = Path(session_dir) / self.scheme_filename

        try:
            from docx import Document
        except ImportError:
            print("生成方案文档需要 python-docx，已跳过")
            return None

        try:
            mapping = self._build_scheme_mapping(filled_data, result)
            document = Document(self.scheme_template_path)
            for paragraph in document.paragraphs:
                self._fill_paragraph(paragraph, mapping)
            document.save(scheme_path)
            print(f"方案文档已生成: {scheme_path}")
            return str(scheme_path)
        except Exception as e:
            print(f"生成方案文档失败: {str(e)}")
            return None

    @staticmethod
    def _fill_paragraph(paragraph, mapping: dict):
        """将段落中的【】占位符替换为实际值"""
        text = paragraph.text
        if "【" not in text:
            return
        for key, value in mapping.items():
            text = text.replace(key, str(value))
        runs = paragraph.runs
        if runs:
            runs[0].text = text
            for run in runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(text)

    def _build_scheme_mapping(self, filled_data: dict, result: dict) -> dict:
        """构造空调方案文档占位符到实际值的映射表"""
        return {
            "【项目区域】": str(filled_data.get("project_region", "")),
            "【站房类型】": str(filled_data.get("station_type", "")),
            "【数据时间段】": f"{datetime.now().year}年",
            "【电价(元/kWh)】": str(filled_data.get("electricity_price", "0.6")),
            "【数据完整率(%)】": f"{result.get('completeness', 0) * 100:.1f}",
            "【总节电量(万kWh)】": f"{result.get('total_saving', 0):.1f}",
            "【冷机节电(万kWh)】": f"{result.get('chiller_saving', 0):.1f}",
            "【水泵节电(万kWh)】": f"{result.get('pump_saving', 0):.1f}",
            "【冷却塔节电(万kWh)】": f"{result.get('tower_saving', 0):.1f}",
            "【综合节能率(%)】": f"{result.get('energy_saving_rate', 0) * 100:.1f}",
            "【节能率范围】": str(result.get("energy_saving_rate_range", "")),
        }

    def _calculate_single_sheet(
        self,
        data: Dict[str, Any],
        excel_path: Path,
        sheet_type: str
    ) -> Dict[str, Any]:
        """计算单个sheet（上限或下限）

        Args:
            data: 填充后的数据字典
            excel_path: Excel 文件路径
            sheet_type: sheet类型 ("upper"/"lower")

        Returns:
            单个sheet的计算结果
        """
        import formulas
        import openpyxl

        # 获取sheet名称
        sheet_name = self.mapper.get_sheet_name(sheet_type)

        # 阶段1: 使用 openpyxl 将数据写入 Excel 文件
        wb = openpyxl.load_workbook(excel_path)

        # 验证sheet是否存在（添加详细错误信息）
        if sheet_name not in wb.sheetnames:
            available_sheets = wb.sheetnames
            wb.close()
            raise ValueError(
                f"Sheet '{sheet_name}' 不存在于Excel模板中。\n"
                f"可用的sheet名称: {available_sheets}\n"
                f"提示: 检查sheet名称是否有多余空格或特殊字符\n"
                f"配置位置: config/cell_mapping.py:SHEET_NAMES['{sheet_type}']"
            )

        sheet = wb[sheet_name]

        # 1. 回填项目地区
        region_cell = self.mapper.get_basic_info_cell("project_region")
        sheet[region_cell] = data["project_region"]

        # 回填电费单价
        electricity_price = data.get("electricity_price")
        if electricity_price is not None:
            price_cell = self.mapper.get_basic_info_cell("electricity_price")
            sheet[price_cell] = electricity_price
        
        # 准备气象数据 - 为每个月份填入对应的湿球温度
        for month_idx in range(12):
            month = month_idx + 1
            wet_bulb_temp = self.weather_service.get_cooling_tower_temp(data["project_region"], month)
            # 在Excel模板中设定特定位置存放月度气象数据，这里假设为Q列开始
            # 实际模板可能需要预设气象数据的位置
            temp_cell = f'Q{month_idx + 6}'  # 从第6行开始放置12个月的气象数据
            sheet[temp_cell] = wet_bulb_temp

        # 2. 回填月度数据
        for month_data in data["monthly_data"]:
            month_idx = month_data["month"] - 1

            for field in self.mapper.get_monthly_fields():
                value = month_data.get(field)
                if value is not None:
                    cell = self.mapper.get_cell("monthly_data", month_idx, field)
                    sheet[cell] = value

        # 3. 保存写入的数据
        wb.save(excel_path)
        wb.close()

        # 阶段2: 使用 formulas 库加载并计算
        xl_model = formulas.ExcelModel().loads(str(excel_path)).finish()
        book_name = list(xl_model.books.keys())[0]
        result_data = xl_model.calculate()

        # 阶段3: 将计算结果写回Excel文件
        wb = openpyxl.load_workbook(excel_path)
        sheet = wb[sheet_name]
        
        # 遍历所有单元格，将公式的计算结果写回
        for row in range(1, sheet.max_row + 1):
            for col in range(1, sheet.max_column + 1):
                cell = sheet.cell(row=row, column=col)
                
                # 如果是公式，获取计算结果并写回
                if cell.data_type == 'f':
                    # 构建cell_id
                    cell_id = f"'[{book_name.lower()}]{sheet_name}'!{cell.coordinate}"
                    
                    # 从formulas结果中获取计算值
                    if cell_id in result_data:
                        value = result_data[cell_id]
                        if hasattr(value, 'value'):
                            calculated_value = value.value[0][0]
                            # 直接写入计算结果，不尝试保留格式
                            cell.value = calculated_value
        
        # 保存写入的计算结果
        wb.save(excel_path)
        wb.close()

        # 阶段4: 提取计算结果
        _excel_error_prefixes = ("#",)  # #DIV/0!、#N/A、#VALUE! 等 Excel 错误值

        def get_result_value(result_field: str) -> float:
            """提取结果单元格的值（P1 改造：失败关闭）

            原实现：单元格缺失/结果为 None 时静默返回 0.0，把计算故障伪装成
            "无节能效果"直接输出给销售报价。现改为任何异常情况显式报错。
            """
            import math

            cell = self.mapper.get_result_cell(result_field)
            cell_id = f"'[{book_name.lower()}]{sheet_name}'!{cell}"

            if cell_id not in result_data:
                raise RuntimeError(
                    f"Excel 计算结果缺少 {sheet_name}!{cell}（字段 {result_field}）——"
                    "公式单元格不存在或映射失配，拒绝静默按 0 处理。"
                    "请核对 config/cell_mapping.py 与公式模板版本。"
                )
            value = result_data[cell_id]
            if hasattr(value, 'value'):
                value = value.value[0][0]
            if value is None:
                raise RuntimeError(
                    f"Excel 公式结果为空（{sheet_name} {cell}，字段 {result_field}）——"
                    "公式未计算或输入缺失，拒绝静默按 0 处理。"
                )
            text = str(value).strip()
            if text.startswith(_excel_error_prefixes) and not _is_number(text):
                raise RuntimeError(
                    f"Excel 公式结果为错误值 {text}（{sheet_name} {cell}，字段 {result_field}）。"
                    "通常是输入数据缺失（如能耗/温差/频率为空）导致公式除零——"
                    "请检查调研表提取与填充结果，而不是静默采用无效数值。"
                )
            num = float(value)
            if not math.isfinite(num):
                raise RuntimeError(
                    f"Excel 公式结果非有限数值（{sheet_name} {cell}，字段 {result_field} = {num}）"
                )
            return num

        def _is_number(text: str) -> bool:
            try:
                float(text)
                return True
            except (TypeError, ValueError):
                return False

        result = {
            "chiller_saving": get_result_value("chiller_saving_total"),
            "pump_saving": (
                get_result_value("cooling_pump_saving_total") +
                get_result_value("freezing_pump_saving_total")
            ),
            "tower_saving": get_result_value("tower_saving_total"),
            "total_saving": get_result_value("total_saving"),
            "saving_rate": get_result_value("saving_rate"),
        }

        return result