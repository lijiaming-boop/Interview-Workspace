from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r'E:\contion simulation\李佳明_简历_Agent开发岗位_精排版.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Mm(16), Mm(16)
sec.left_margin, sec.right_margin = Mm(18), Mm(18)
sec.header_distance, sec.footer_distance = Mm(7), Mm(8)

def font(obj, size=10.5, bold=False, color='252525'):
    obj.font.name = 'Calibri'
    obj.font.size = Pt(size)
    obj.font.bold = bold
    obj.font.color.rgb = RGBColor.from_string(color)
    rp = obj._element.get_or_add_rPr()
    rf = rp.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rp.insert(0, rf)
    rf.set(qn('w:eastAsia'), '微软雅黑')
    rf.set(qn('w:ascii'), 'Calibri')
    rf.set(qn('w:hAnsi'), 'Calibri')

normal = doc.styles['Normal']
font(normal)
pf = normal.paragraph_format
pf.space_after, pf.line_spacing = Pt(4), Pt(15.5)
pf.widow_control = True
for s, size in [('Title',24), ('Heading 1',13), ('Heading 2',11.5)]:
    font(doc.styles[s], size, True, '000000')
    doc.styles[s].paragraph_format.keep_with_next = True
    doc.styles[s].paragraph_format.space_before = Pt(10 if s == 'Heading 1' else 5)
    doc.styles[s].paragraph_format.space_after = Pt(5)
    doc.styles[s].paragraph_format.line_spacing = Pt(30 if s == 'Title' else 17)

# Clear template decoration and CJK document-grid overrides.
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    if style.type == 1:
        prop = style.element.get_or_add_pPr()
        snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); prop.append(snap)
for grid in list(sec._sectPr.findall(qn('w:docGrid'))):
    sec._sectPr.remove(grid)

def p(text='', lead=False, size=10.5, after=4):
    x=doc.add_paragraph()
    x.paragraph_format.space_after=Pt(after)
    x.paragraph_format.keep_together=True
    if lead and '：' in text:
        a,b=text.split('：',1)
        font(x.add_run(a+'：'),size,True)
        font(x.add_run(b),size)
    else: font(x.add_run(text),size)
    return x

def heading(text, page=False):
    x=doc.add_paragraph(text,'Heading 1')
    x.paragraph_format.page_break_before=page
    if page: x.paragraph_format.space_before=Pt(0)
    return x

def dated(title,date):
    x=doc.add_paragraph(style='Heading 2')
    x.paragraph_format.tab_stops.add_tab_stop(Mm(174),WD_TAB_ALIGNMENT.RIGHT)
    font(x.add_run(title),11.5,True,'000000')
    font(x.add_run('\t'+date),10,False,'555555')
    return x

def meta(text):
    x=p(text,size=9.5,after=5)
    for r in x.runs: r.font.color.rgb=RGBColor.from_string('606060')
    x.paragraph_format.keep_with_next=True

name=doc.add_paragraph('李佳明','Title')
name.alignment=WD_ALIGN_PARAGRAPH.CENTER
name.paragraph_format.space_before=Pt(0)
name.paragraph_format.space_after=Pt(5)
x=p('求职意向  Agent 开发工程师 / AI 应用开发    |    2027 届本科毕业生',size=11,after=5)
x.alignment=WD_ALIGN_PARAGRAPH.CENTER
x=p('18995488153  |  18995488153@163.com',size=10,after=3)
x.alignment=WD_ALIGN_PARAGRAPH.CENTER
x=p('男 · 21 岁 · 共青团员    |    现居辽宁沈阳 · 生源地宁夏',size=9.5,after=5)
x.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in x.runs:r.font.color.rgb=RGBColor.from_string('606060')

heading('教育背景')
dated('东北大学  ·  计算机科学与技术', '2023.09 – 2027.06')
p('工学学士（在读）  |  985 / 211 / 双一流  |  CET-4、CET-6 已通过',after=3)
p('主修课程：数据结构与算法、操作系统、计算机网络、数据库原理、软件工程。',lead=True)

heading('专业技能')
p('Agent 与大模型：具备 LLM API 对接、Prompt 设计、ReAct 工作流、Skill 脚本工具化及 JSON Schema 约束经验；熟悉工具调用、人工介入、轨迹记录与异常回退。',lead=True)
p('Python 与后端：熟悉 Python、FastAPI、REST / SSE、SQLAlchemy 及 SQLite / PostgreSQL；使用 pandas、NumPy、SciPy、openpyxl 完成数据处理、建模与结果分析。',lead=True)
p('C++ 与系统开发：熟悉 C++17、STL、RAII、多线程、有界队列与 SPSC 环形缓冲区；具备 ROS 2 / DDS、RPC 服务集成、异步 IO 和性能优化经验。',lead=True)
p('工程化与算法：使用 Git、CMake、GDB、pytest、GoogleTest、Docker 与 CI；具备 SLSQP 约束优化、闭环仿真、交叉验证及物理安全校验经验。',lead=True)

heading('实习经历')
dated('蘑菇物联  ·  AI 算法开发实习生', '2026.04 – 2026.09')
meta('暖通节能工具开发与制冷站智能优化')
x=doc.add_paragraph('暖通工具测算 Agent','Heading 2')
p('业务流程：面向暖通节能售前场景，参与调研表提取、质检、缺失填充、公式计算与结论生成流程开发，将 Excel 测算能力封装为可复用工具。',lead=True)
p('Agent 编排：构建 ReAct 循环与 LLM / 规则双模式决策，配置 Prompt、JSON 约束、工具白名单、前置依赖及质量门禁；数据不足时转人工处理。',lead=True)
p('服务与产出：通过 FastAPI / SSE 提供测算接口与过程反馈，支持会话持久化、轨迹记录和断点恢复；规则模式下 9 份样例流程决策全部正确，低质量数据拦截 2/2。',lead=True)
doc.add_paragraph('制冷站功率预测与闭环联合寻优','Heading 2')
p('建模与优化：清洗历史工况，建立冷却塔、冷却泵、主机功率与逼近温差子模型；结合塔组合枚举和 SLSQP 调频，在流量、温度及设备容量约束下输出运行建议。',lead=True)
p('验证成果：形成六项目验证报告；B/C 测试集总功率 R² 为 0.969 / 0.977，可行工况相对实测基线的离线估计中位节能率为 6.74% / 3.32%。',lead=True)
p('智能推理探索：参与 CoT 数据工厂与物理决策链数据集构建，设计物理安全校验及仿真回退流程，为后续小模型推理提供数据与验证基础。',lead=True)

heading('个人项目',page=True)
dated('自动驾驶数据缓存与回流系统','2026.08 – 至今')
meta('C++17 / ROS 2 Jazzy / DDS / zstd / OpenCV / PCL')
p('项目概述：独立开发相机与激光雷达事件触发数据系统，打通持续缓存、事件前后窗口截取、压缩落盘、图像/点云转换、远端回传与回读校验流程。',lead=True)
p('并发与性能：设计分传感器缓存、近似时间同步与配对索引；将序列化、压缩和 IO 移至后台线程，实现热路径摊还 O(1)。局部基准测试中，PairIndex / DataBuffer 热点改善约 980 / 385 倍。',lead=True)
p('可靠性设计：实现传感器时间窗口与墙钟兜底、看门狗和磁盘保留策略；用文件标记管理存储/上传状态，通过临时文件、fsync、原子重命名和完整性校验保障落盘一致性。',lead=True)
p('接口与回传：设计 UploadStore ROS 2 RPC，以 BEGIN / FILE_CHUNK / END 分块传输；结合 transfer_id、SHA-256 清单、指数退避、幂等触发和生命周期状态回传处理异常。',lead=True)
p('测试交付：实现 record_reader 校验/导出、模拟接收端与端到端冒烟测试；项目记录覆盖 48 个单测、1,013 条集成断言，并接入 CI 编译、格式检查及 ASan / UBSan。',lead=True)

dated('数字人口型同步与实时 AI 会话 SDK','2026.01 – 2026.03')
meta('C++17 / ncnn / FFmpeg / JSON / SSE / RTMP / RTSP')
p('项目概述：独立开发音视频实时生成 SDK，集成模型推理、音频处理、画面合成、流式会话与推流，将音频和人脸图像生成口型同步视频。',lead=True)
p('大模型集成：对接 OpenAI 兼容大模型服务，解析 JSON / SSE 增量响应并按句驱动流式 TTS；通过显式会话状态机、统一超时和媒体错误终态管理服务生命周期。',lead=True)
p('推理与调度：建立 VAD、Mel 特征提取和音频时钟驱动的帧调度链路，推理调用由约 100 次/秒降至 25 次/秒；统一 ncnn 张量布局并引入 SCRFD，修复推理和侧脸检测问题。',lead=True)
p('并发与渲染：使用线程安全队列和 SPSC 环形缓冲区处理背压；以嘴部 ROI 合成、掩码融合和静态人脸缓存，将渲染耗时由 113 ms 降至 14.3 ms（约 7.9 倍提升）。',lead=True)
p('推流与验证：基于 FFmpeg 实现 H.264 / AAC 编码回退和 RTMP / RTSP 发布；建立 5 级分层测试体系，项目记录累计 450+ 用例通过。',lead=True)

heading('校园经历与荣誉')
p('班级宣传委员：2023.09 至今，负责活动宣传、主题班会策划及通知材料整理。',lead=True)
p('荣誉奖项：军训先进个人（2023.09）。',lead=True)
heading('自我评价')
p('关注 AI 能力与业务流程的结合，能围绕数据处理、工具编排、服务集成和验证交付推进工作；重视结果可解释性、异常回退及团队协作。遵守纪律，认真负责。')

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
font(footer.add_run('李佳明  ·  Agent 开发工程师    '),8,False,'777777')
for value in ['PAGE','NUMPAGES']:
    r=footer.add_run();font(r,8,False,'777777')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),value);r._r.addnext(field)
    if value=='PAGE':font(footer.add_run(' / '),8,False,'777777')
doc.core_properties.title='李佳明 Agent开发岗位简历'
doc.core_properties.author='李佳明'
doc.core_properties.subject='AI应用开发 暖通测算Agent 智能优化'
doc.save(OUT)
print(OUT)
