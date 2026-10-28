from copy import deepcopy
from pathlib import Path

from docx import Document


SOURCE = Path(r"E:\contion simulation\李佳明_简历_Agent开发岗位_优化版.docx")
OUTPUT = Path(r"E:\contion simulation\李佳明_简历_Agent开发岗位_实习个人项目版.docx")


def replace_paragraph(paragraph, text: str) -> None:
    """Replace visible text while retaining the paragraph style and first-run formatting."""
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def replace_cell(cell, text: str) -> None:
    replace_paragraph(cell.paragraphs[0], text)


doc = Document(str(SOURCE))

# Align the headline and skill profile with the requested Agent-oriented role.
replace_paragraph(
    doc.paragraphs[8],
    "Python 与 AI 应用：熟悉 Python 3.10+、pandas、NumPy、SciPy、openpyxl；具备历史工况清洗、特征构造、模型拟合、交叉验证、指标分析与结果整理经验。",
)
replace_paragraph(
    doc.paragraphs[9],
    "大模型与智能推理：具备将物理状态结构化为 6 步 CoT 决策链、构建训练/验证数据集、设计小模型推理与 Prompt 约束的项目经验；了解 Agent 工作流与安全回退思路。",
)
replace_paragraph(
    doc.paragraphs[10],
    "建模与优化：掌握幂律模型、非线性回归、SLSQP 约束优化、闭环迭代、塔组合枚举、数据域包络与物理约束校验；能够区分模型基线、实测基线与离线节能估计。",
)
replace_paragraph(
    doc.paragraphs[11],
    "C++ 与 ROS 2：熟悉 C++17、STL、RAII、多线程同步、ROS 2 Topic / Service、DDS 与 Launch；具备传感器数据缓存、事件触发、异步写盘和 RPC 回传开发经验。",
)
replace_paragraph(
    doc.paragraphs[12],
    "后端与数据工程：熟悉 JSON / CSV / Excel 数据处理、SQLite / PostgreSQL / Redis 基础；了解 CMake、GCC、GDB、pytest、GoogleTest、Docker 与 CI，重视可测试性和故障可追溯。",
)
replace_paragraph(
    doc.paragraphs[13],
    "工程化能力：能够拆解业务问题，完成数据处理、模型/规则设计、服务集成、测试验证与结果解释；关注性能边界、异常回退和交付文档。",
)

# The internship is the jianmo work at 蘑菇物联; datacache moves to personal projects below.
replace_cell(
    doc.tables[1].rows[1].cells[0],
    "蘑菇物联  ｜  制冷站功率预测与联合寻优  ｜  AI 算法开发实习生  ｜  2026.04 - 2026.09",
)
replace_paragraph(
    doc.paragraphs[16],
    "数据建模与验证：参与制冷站历史工况数据清洗、缓存化和字段统一，建立冷却塔散热、冷却泵功率、主机功率与逼近温差子模型；采用 70%/30% 划分和时间分块交叉验证评估泛化。",
)
replace_paragraph(
    doc.paragraphs[17],
    "闭环联合寻优：将实测主机功率 ChPower、冷却泵功率 CWPPower 纳入目标函数，完成塔组合枚举 + SLSQP 连续调频；加入流量/风量包络、温差、出水温度与单塔容量约束，输出泵/塔频率和净功率建议。",
)
replace_paragraph(
    doc.paragraphs[18],
    "产出与质量：形成六项目验证、可行率与节电口径报告；B/C 测试集总功率 R² 达 0.969 / 0.977，B/C 可行工况相对实测基线中位节能约 6.74% / 3.32%；参与构建 CoT 数据工厂、物理决策链数据集与安全校验/回退推理方案。",
)

# Replace the former Agent personal project with the actual datacache repository.
replace_paragraph(
    doc.paragraphs[20],
    "ROS 2 自动驾驶数据缓存与回流系统（个人项目）｜C++17 / ROS 2 Jazzy / zstd / OpenCV / PCL / DDS｜2026.08 - 至今",
)
replace_paragraph(
    doc.paragraphs[21],
    "项目概述：独立开发车载相机 / 激光雷达事件触发数据缓存系统，完成持续缓存、碰撞/急刹触发、pre/post 窗口截取、压缩落盘、图像/点云转换、远端回传与回读校验闭环。",
)
replace_paragraph(
    doc.paragraphs[22],
    "并发与热路径：设计按传感器分队列的 DataBuffer、ApproximateSynchronizer 近似配对、PairIndex 同步账本和 EventMonitor 调度器；回调线程只入队，RawStorageWorker 后台完成 CDR 序列化、zstd 压缩和 IO，实现热路径摊还 O(1)。",
)
replace_paragraph(
    doc.paragraphs[23],
    "可靠性与状态机：实现传感器时间域窗口 + 墙钟兜底、SensorWatchdog、写前磁盘检查、容量/天数保留、.pending/.complete/.failed/.uploaded 状态标记，以及 .tmp + fsync + rename 原子落盘和完整性校验。",
)
replace_paragraph(
    doc.paragraphs[24],
    "服务集成与回传：实现 UploadWorker 与 upload_receiver_node，通过 UploadStore ROS 2 RPC 按 BEGIN / FILE_CHUNK / END 分块传输；加入 transfer_id、SHA-256 清单、重试退避、幂等触发与 RECEIVED → ACCEPTED → STORED → UPLOADED 状态回传。",
)
replace_paragraph(
    doc.paragraphs[25],
    "测试与性能：补齐 record_reader 回读/导出、mock 接收端、端到端 smoke 和 GitHub Actions 编译/单测/clang-format/ASan/UBSan 门禁；8 个 gtest + 1 个 pytest 共 48 个单测，集成测试 1,013 条断言，PairIndex / DataBuffer 热路径优化约 980x / 385x。",
)

replace_paragraph(
    doc.paragraphs[38],
    "Agent 应用与工程落地：能够从数据与业务问题出发，完成数据处理、模型/规则设计、约束优化、服务集成、测试与结果解释；能将复杂决策过程结构化，并关注安全护栏、异常回退与用户可理解性。",
)
replace_paragraph(
    doc.paragraphs[39],
    "工程协作与综合素质：实习期间参与制冷站功率预测与联合寻优并形成验证报告；个人项目中独立完成 ROS 2 数据缓存回流系统，具备问题拆解、进度推进和稳定交付能力；共青团员，遵守纪律，品行端正。",
)

doc.save(str(OUTPUT))
print(OUTPUT)
